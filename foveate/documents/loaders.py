"""Loaders turn files into page texts; `load` assembles a `Document`."""

from __future__ import annotations

import abc
import dataclasses
import html.parser
import io
import pathlib
from typing import ClassVar, Literal

from foveate import errors
from foveate.documents import document as document_lib
from foveate.documents import page as page_lib
from foveate.internals import registry
from foveate.partition import base as partition_base
from foveate.tokenizers import base as tokenizer_base

SUFFIXES = {
    ".pdf": "pdf",
    ".txt": "text",
    ".text": "text",
    ".log": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".html": "html",
    ".htm": "html",
    ".docx": "docx",
}
FORM_FEED = "\f"
DEFAULT_PAGE_TOKENS = 500


def paginate(
    text: str, tokenizer: tokenizer_base.Tokenizer, page_tokens: int
) -> list[str]:
    """Splits flowing text into pseudo-pages of about `page_tokens` tokens.

    Form feeds are honoured as real page breaks; otherwise the lossless
    recursive partitioner cuts at paragraph and sentence boundaries.
    """
    if FORM_FEED in text:
        return text.split(FORM_FEED)
    if not text.strip():
        return []
    splitter = partition_base.RecursivePartitioner(page_tokens)
    return [p.text for p in splitter.split(text, tokenizer)]


@dataclasses.dataclass(frozen=True)
class Loader(abc.ABC):
    """Extracts page texts from raw bytes.

    Subclasses are dataclasses whose fields are their options.
    """

    registry: ClassVar[registry.Registry[type[Loader]]] = registry.Registry(
        "document format"
    )
    page_tokens: int = DEFAULT_PAGE_TOKENS

    def __post_init__(self) -> None:
        if self.page_tokens < 1:
            raise errors.ConfigError("page_tokens must be >= 1")

    @abc.abstractmethod
    def extract(
        self, data: bytes, tokenizer: tokenizer_base.Tokenizer
    ) -> list[str]:
        """Returns the text of each page, in order.

        Raises:
            ValidationError: If the data cannot be parsed.
            ConfigError: If an optional dependency is missing.
        """


def decode(data: bytes) -> str:
    """Decodes UTF-8 text (BOM tolerated) and normalises line endings to LF.

    Raises:
        ValidationError: If the bytes are not valid UTF-8.
    """
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise errors.ValidationError(f"not valid UTF-8: {exc}") from exc
    return text.replace("\r\n", "\n").replace("\r", "\n")


@Loader.registry.register("text")
@dataclasses.dataclass(frozen=True)
class TextLoader(Loader):
    """Plain text; form feeds mark pages, else pseudo-pages are cut."""

    def extract(
        self, data: bytes, tokenizer: tokenizer_base.Tokenizer
    ) -> list[str]:
        return paginate(decode(data), tokenizer, self.page_tokens)


@Loader.registry.register("markdown")
@dataclasses.dataclass(frozen=True)
class MarkdownLoader(TextLoader):
    """Markdown, paginated like text (headings are detected per page)."""


class TextExtractor(html.parser.HTMLParser):
    """Collects visible text from HTML, one line per block element."""

    BLOCKS = frozenset(
        [
            "p",
            "div",
            "br",
            "li",
            "ul",
            "ol",
            "tr",
            "table",
            "section",
            "article",
            "header",
            "footer",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "pre",
            "blockquote",
        ]
    )
    SKIP = frozenset({"script", "style", "noscript", "head"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skipping = 0

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag in self.SKIP:
            self.skipping += 1
        if tag in self.BLOCKS:
            self.parts.append("\n")
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("# ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skipping:
            self.parts.append(data)


@Loader.registry.register("html")
@dataclasses.dataclass(frozen=True)
class HtmlLoader(Loader):
    """HTML reduced to visible text; headings become Markdown `#` lines."""

    def extract(
        self, data: bytes, tokenizer: tokenizer_base.Tokenizer
    ) -> list[str]:
        extractor = TextExtractor()
        extractor.feed(decode(data))
        lines = [line.strip() for line in "".join(extractor.parts).splitlines()]
        text = "\n".join(line for line in lines if line)
        return paginate(text, tokenizer, self.page_tokens)


@Loader.registry.register("pdf")
@dataclasses.dataclass(frozen=True)
class PdfLoader(Loader):
    """PDF via `pypdf` (install the `pdf` extra). Real page boundaries.

    Attributes:
        layout: Use pypdf's layout-preserving extraction (better tables,
            slower).
    """

    layout: bool = False

    def extract(
        self, data: bytes, tokenizer: tokenizer_base.Tokenizer
    ) -> list[str]:
        try:
            import pypdf
        except ImportError as exc:
            raise errors.ConfigError(
                "PDF support needs pypdf: pip install 'foveate[pdf]'"
            ) from exc
        try:
            reader = pypdf.PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise errors.ValidationError("PDF is password protected")
            mode: Literal["plain", "layout"] = (
                "layout" if self.layout else "plain"
            )
            return [
                page.extract_text(extraction_mode=mode) or ""
                for page in reader.pages
            ]
        except errors.FoveateError:
            raise
        except Exception as exc:
            raise errors.ValidationError(f"cannot read PDF: {exc}") from exc


@Loader.registry.register("docx")
@dataclasses.dataclass(frozen=True)
class DocxLoader(Loader):
    """Word documents via `python-docx` (install the `docx` extra).

    Word has no stable page boundaries, so text is cut into pseudo-pages.
    """

    def extract(
        self, data: bytes, tokenizer: tokenizer_base.Tokenizer
    ) -> list[str]:
        try:
            import docx
        except ImportError as exc:
            raise errors.ConfigError(
                "DOCX support needs python-docx: pip install 'foveate[docx]'"
            ) from exc
        try:
            parsed = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise errors.ValidationError(f"cannot read DOCX: {exc}") from exc
        lines = []
        for paragraph in parsed.paragraphs:
            style = (paragraph.style.name or "") if paragraph.style else ""
            prefix = "# " if style.startswith("Heading") else ""
            if paragraph.text.strip():
                lines.append(prefix + paragraph.text)
        return paginate("\n\n".join(lines), tokenizer, self.page_tokens)


def load(
    source: str | pathlib.Path | bytes,
    *,
    format: str | None,
    doc_id: str | None,
    tokenizer: tokenizer_base.Tokenizer | None,
    options: dict[str, object],
) -> document_lib.Document:
    """Builds a `Document`; see `Document.load` for the contract."""
    tokenizer = tokenizer or tokenizer_base.default_tokenizer()
    if isinstance(source, bytes):
        if format is None:
            raise errors.ConfigError("loading bytes needs format=...")
        data, origin, stem = source, "", doc_id or "document"
    else:
        path = pathlib.Path(source)
        try:
            data = path.read_bytes()
        except OSError as exc:
            raise errors.ValidationError(f"cannot read {path}: {exc}") from exc
        format = format or SUFFIXES.get(path.suffix.lower())
        if format is None:
            raise errors.ConfigError(
                f"cannot infer format of {path.name}; pass format=..."
            )
        origin, stem = str(path), doc_id or path.stem
    loader_cls = Loader.registry.get(format)
    try:
        loader = loader_cls(**options)  # type: ignore[arg-type]
    except TypeError as exc:
        raise errors.ConfigError(
            f"invalid options for {format!r} loader: {exc}"
        ) from exc
    texts = loader.extract(data, tokenizer)
    if not texts:
        raise errors.ValidationError("the document has no extractable text")
    pages = tuple(
        page_lib.Page(
            number=index,
            text=text,
            tokens=tokenizer.count(text),
            headings=page_lib.detect_headings(text),
        )
        for index, text in enumerate(texts, start=1)
    )
    empty = sum(1 for p in pages if not p.text.strip())
    return document_lib.Document(
        id=stem,
        pages=pages,
        title=stem,
        source=origin,
        metadata={"format": format, "empty_pages": str(empty)},
    )
