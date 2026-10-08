import pytest
from hypothesis import given
from hypothesis import strategies as st

from foveate import Message, Role, errors
from foveate.documents import (
    Document,
    Loader,
    detect_headings,
    parse_pages,
    unique_ids,
)
from foveate.documents import page as page_lib
from foveate.tokenizers import HeuristicTokenizer

TOK = HeuristicTokenizer()


def make_pdf(pages: int, blank: tuple[int, ...] = ()) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for number in range(1, pages + 1):
        pdf.add_page()
        if number in blank:
            continue
        pdf.set_font("Helvetica", style="B", size=14)
        pdf.cell(
            0, 10, f"SECTION {number} OVERVIEW", new_x="LMARGIN", new_y="NEXT"
        )
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(
            0, 6, f"Marker PAGEMARK{number:03d}. " + "Filler text. " * 30
        )
    return bytes(pdf.output())


class TestPageSpecs:
    def test_forms(self):
        avail = list(range(1, 21))
        assert parse_pages("3", avail) == (3,)
        assert parse_pages("10-12, 5,5", avail) == (5, 10, 11, 12)
        assert parse_pages(7, avail) == (7,)
        assert parse_pages(range(2, 5), avail) == (2, 3, 4)

    @pytest.mark.parametrize("bad", ["", "a", "3-", "9-2", "0", "99", "1,,x"])
    def test_invalid(self, bad):
        with pytest.raises(errors.ValidationError):
            parse_pages(bad, list(range(1, 21)))

    @given(st.sets(st.integers(1, 50), min_size=1, max_size=10))
    def test_roundtrip(self, numbers):
        spec = ",".join(str(n) for n in numbers)
        assert parse_pages(spec, list(range(1, 51))) == tuple(sorted(numbers))


def test_heading_heuristics():
    text = "\n".join(
        [
            "# Markdown title",
            "2.1 Scope of work",
            "ITEM 7. MANAGEMENT DISCUSSION",
            "This is an ordinary sentence that ends here.",
            "SHORT",
            "12345",
            "A list item,",
        ]
    )
    found = detect_headings(text)
    assert "Markdown title" in found and "2.1 Scope of work" in found
    assert "ITEM 7. MANAGEMENT DISCUSSION" in found
    assert not any("ordinary" in h or h == "12345" for h in found)


class TestDocument:
    def doc(self, n=10):
        pages = tuple(
            page_lib.Page(i, f"text of page {i}", 4, (f"H{i}",))
            for i in range(1, n + 1)
        )
        return Document("report", pages, title="Report")

    def test_select_keeps_original_numbers_and_around_clips(self):
        doc = self.doc()
        sub = doc.select("2-3,9")
        assert sub.numbers == [2, 3, 9] and sub.page(9).text == "text of page 9"
        assert doc.around(1, 2).numbers == [1, 2, 3]
        assert doc.around(10, 1).numbers == [9, 10]
        with pytest.raises(errors.ValidationError):
            doc.around(11)
        with pytest.raises(errors.ValidationError):
            sub.page(5)

    def test_outline_text_and_messages(self):
        doc = self.doc(3)
        assert [(h.title, h.page) for h in doc.outline()] == [
            ("H1", 1),
            ("H2", 2),
            ("H3", 3),
        ]
        assert doc.text(markers=True).startswith("[report p.1]\ntext of page 1")
        messages = doc.to_messages()
        assert messages[0] == Message(Role.USER, "[report p.1]\ntext of page 1")
        assert doc.token_count == 12

    def test_validation(self):
        with pytest.raises(errors.ValidationError):
            Document("", ())
        with pytest.raises(errors.ValidationError):
            Document("d", (page_lib.Page(2, "a", 1), page_lib.Page(1, "b", 1)))
        with pytest.raises(errors.ValidationError):
            page_lib.Page(0, "x", 1)
        with pytest.raises(errors.ValidationError):
            unique_ids([self.doc(), self.doc()])


class TestLoaders:
    def test_pdf_has_real_pages_and_headings(self, tmp_path):
        path = tmp_path / "filing.pdf"
        path.write_bytes(make_pdf(12, blank=(5,)))
        doc = Document.load(path, tokenizer=TOK)
        assert doc.id == "filing" and len(doc.pages) == 12
        assert (
            "PAGEMARK007" in doc.page(7).text
            and "PAGEMARK007" not in doc.page(8).text
        )
        assert (
            doc.metadata["empty_pages"] == "1" and not doc.page(5).text.strip()
        )
        assert any("SECTION 3" in h for h in doc.page(3).headings)
        assert doc.select("7").text().count("PAGEMARK") == 1
        assert [m.content[:14] for m in doc.select("1-2").to_messages()] == [
            "[filing p.1]\nS",
            "[filing p.2]\nS",
        ]

    def test_pdf_from_bytes_and_layout_option(self):
        doc = Document.load(make_pdf(2), format="pdf", doc_id="b", layout=True)
        assert doc.id == "b" and len(doc.pages) == 2

    def test_corrupt_and_missing_pdf(self, tmp_path):
        with pytest.raises(errors.ValidationError):
            Document.load(b"%PDF-1.4 not really", format="pdf")
        with pytest.raises(errors.ValidationError):
            Document.load(tmp_path / "missing.pdf")

    def test_text_pagination_is_lossless_and_form_feed_wins(self, tmp_path):
        body = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(40))
        path = tmp_path / "big.txt"
        path.write_text(body, encoding="utf-8")
        doc = Document.load(path, tokenizer=TOK, page_tokens=200)
        assert len(doc.pages) > 5 and all(p.tokens <= 200 for p in doc.pages)
        assert "".join(p.text for p in doc.pages) == body
        paged = Document.load(b"one\fTwo\fthree", format="text")
        assert [p.text for p in paged.pages] == ["one", "Two", "three"]

    def test_markdown_and_html(self, tmp_path):
        md = Document.load(b"# Title\n\nHello world.", format="markdown")
        assert md.outline()[0].title == "Title"
        html = Document.load(
            b"<html><head><style>x{}</style></head><body><h1>Big Title</h1>"
            b"<p>First para.</p><script>bad()</script><p>Second.</p></body></html>",
            format="html",
        )
        assert "bad()" not in html.text() and "x{}" not in html.text()
        assert (
            html.outline()[0].title == "Big Title" and "Second." in html.text()
        )

    def test_docx(self, tmp_path):
        import docx

        d = docx.Document()
        d.add_heading("Quarterly Plan", level=1)
        d.add_paragraph("We will ship the feature in March.")
        path = tmp_path / "plan.docx"
        d.save(path)
        doc = Document.load(path)
        assert (
            "ship the feature" in doc.text()
            and doc.outline()[0].title == "Quarterly Plan"
        )

    def test_errors(self, tmp_path):
        with pytest.raises(errors.ConfigError):
            Document.load(b"abc")  # bytes need a format
        weird = tmp_path / "x.unknown"
        weird.write_bytes(b"hi")
        with pytest.raises(errors.ConfigError):
            Document.load(weird)
        with pytest.raises(errors.ConfigError):
            Document.load(b"hi", format="text", page_tokens=0)
        with pytest.raises(errors.ConfigError):
            Document.load(b"hi", format="text", bogus=1)
        with pytest.raises(errors.ValidationError):
            Document.load(b"\xff\xfe\x00bad", format="text")
        with pytest.raises(errors.ValidationError):
            Document.load(b"   ", format="text")
        with pytest.raises(errors.ConfigError):
            Document.load(b"x", format="nope")

    def test_registry_names(self):
        assert {"pdf", "text", "markdown", "html", "docx"} <= set(
            Loader.registry.names()
        )
